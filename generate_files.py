import docx
from fpdf import FPDF

content = """Relato Confidencial Operacional
O policial Marcos Silva, CPF 123.456.789-00, lotado em Belo Horizonte, reportou que a viatura 1002 precisa de manutenção na embreagem de forma urgente. 
O REDS associado ao último incidente é 2026-123456789-001. Acreditamos que a solução definitiva é o plano de substituição da frota antiga aprovada para a divisão do plantão.
Sugiro ainda enviar um e-mail para marcos.silva@pc.mg.gov.br copiando o telefone (31) 98765-4321."""

# DOCX
doc = docx.Document()
doc.add_paragraph(content)
doc.save('test_report.docx')

# PDF
class PDF(FPDF):
    def header(self):
        self.set_font("Arial", 'B', 15)
        self.cell(0, 10, 'Tira-Voz: Relatorio Experimental', border=False, ln=True, align='C')
        self.ln(10)

pdf = PDF()
pdf.add_page()
pdf.set_font("Arial", size=12)
pdf.multi_cell(0, 10, content)
pdf.output("test_report.pdf")

# MD
with open('test_report.md', 'w') as f:
    f.write(content)

print("Files generated successfully!")
